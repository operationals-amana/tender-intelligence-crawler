"""Seed the association directory from the Master Directory workbook.

The sheet has three columns: association name, a description that doubles as the
category, and a contact number. Descriptions come in three shapes:

    "Kadin Indonesia ALB Member"
    "Financial & Banking Association"
    "Digital Economy - Cluster: Asuransi"

so the category is derived from the description and the digital-economy rows
additionally carry a sub-cluster.
"""
import os
import re

from openpyxl import load_workbook

from api.seed.paths import resolve_data_file
from crawler.models import Association

BLANKS = {"", "-", "n/a", "na", "#n/a", "none", "null"}


def _clean(value):
    if value is None:
        return None
    text = str(value).strip()
    if text.lower() in BLANKS:
        return None
    return text


def _directory_path() -> str:
    return resolve_data_file("Master Directory_ Associations & Phonebook (1).xlsx")


def _categorize(description: str):
    """Return ``(category, sub_cluster)`` for a directory description."""
    if not description:
        return "Uncategorized", None

    lowered = description.lower()

    if "kadin" in lowered:
        return "Kadin Indonesia ALB Member", None

    if "digital economy" in lowered:
        # "Digital Economy - Cluster: Asuransi" -> sub-cluster "Asuransi".
        match = re.search(r"cluster\s*:\s*(.+)$", description, re.IGNORECASE)
        sub_cluster = match.group(1).strip() if match else None
        return "Digital Economy", sub_cluster

    if "financial" in lowered or "banking" in lowered:
        return "Financial & Banking", None

    return description, None


def _abbreviation(name: str):
    """Pull the parenthesised acronym out of a name, when there is one."""
    match = re.search(r"\(([^)]{2,15})\)\s*$", name)
    return match.group(1).strip() if match else None


def seed_associations(db):
    path = _directory_path()
    if not os.path.exists(path):
        print(f"  Directory not found at {path}, skipping association seed")
        return

    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    existing = {(a.name, a.category) for a in db.query(Association).all()}
    added = 0

    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0:  # header
            continue

        name = _clean(row[0]) if row else None
        if not name:
            continue

        description = _clean(row[1]) if len(row) > 1 else None
        contact = _clean(row[2]) if len(row) > 2 else None
        category, sub_cluster = _categorize(description)

        if (name, category) in existing:
            continue
        existing.add((name, category))

        db.add(
            Association(
                name=name,
                abbreviation=_abbreviation(name),
                category=category,
                sub_cluster=sub_cluster,
                contact=contact,
                description=description,
            )
        )
        added += 1

    db.flush()
    print(f"  Seeded {added} associations")
